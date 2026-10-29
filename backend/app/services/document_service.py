import os
import uuid
import logging
from pathlib import Path
from typing import Tuple
from docx import Document as DocxDocument
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from app.config import settings
from app.services.llm_service import llm_service
from app.services.rag_service import rag_service

logger = logging.getLogger("document_service")

class DocumentService:
    def create_document(self, title: str, format_type: str, content: str) -> Tuple[str, str]:
        """
        Generates a new downloadable document (docx, pdf, txt, or md) and saves it to generated/ directory.
        Returns (filename, filepath).
        """
        safe_title = "".join(c for c in title if c.isalnum() or c in (" ", "_", "-")).rstrip()
        safe_title = safe_title.replace(" ", "_") or "Document"
        unique_suffix = uuid.uuid4().hex[:6]
        fmt = format_type.lower().lstrip(".")

        filename = f"{safe_title}_{unique_suffix}.{fmt}"
        output_path = settings.GENERATED_DIR / filename

        if fmt == "docx":
            doc = DocxDocument()
            doc.add_heading(title, level=0)
            for line in content.split("\n"):
                stripped = line.strip()
                if stripped.startswith("## "):
                    doc.add_heading(stripped[3:], level=2)
                elif stripped.startswith("# "):
                    doc.add_heading(stripped[2:], level=1)
                elif stripped.startswith("- ") or stripped.startswith("* "):
                    doc.add_paragraph(stripped[2:], style='List Bullet')
                elif stripped:
                    doc.add_paragraph(stripped)
            doc.save(str(output_path))

        elif fmt == "pdf":
            doc = SimpleDocTemplate(str(output_path), pagesize=letter)
            styles = getSampleStyleSheet()
            flowables = []
            
            # Title
            flowables.append(Paragraph(title, styles['Title']))
            flowables.append(Spacer(1, 14))

            for line in content.split("\n"):
                stripped = line.strip()
                if stripped:
                    # Clean tags that could break XML parsing in ReportLab
                    safe_line = stripped.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                    flowables.append(Paragraph(safe_line, styles['Normal']))
                    flowables.append(Spacer(1, 6))

            doc.build(flowables)

        elif fmt in ["txt", "md"]:
            with open(str(output_path), "w", encoding="utf-8") as f:
                f.write(f"# {title}\n\n" if fmt == "md" else f"{title}\n\n")
                f.write(content)
        else:
            raise ValueError(f"Unsupported format: {format_type}")

        logger.info(f"Created new document: {output_path}")
        return filename, str(output_path)

    def edit_document(
        self,
        original_path: str,
        file_type: str,
        instructions: str
    ) -> Tuple[str, str, str]:
        """
        Reads the original document, applies requested changes via LLM service,
        and saves a brand new version to generated/ without modifying the original.
        Returns (new_filename, new_filepath, summary_of_changes).
        """
        orig_file = Path(original_path)
        if not orig_file.exists():
            raise FileNotFoundError(f"Original file not found: {original_path}")

        # Extract current text from original document
        current_text, _ = rag_service.extract_text(original_path, file_type)

        # Ask LLM to apply modifications
        prompt = (
            f"You are a professional document editor.\n"
            f"USER INSTRUCTIONS: {instructions}\n\n"
            f"--- ORIGINAL DOCUMENT CONTENT ---\n"
            f"{current_text}\n"
            f"--- END ORIGINAL CONTENT ---\n\n"
            f"Please output the complete revised document content incorporating the user's instructions. "
            f"Do not include meta commentary or preface; output the revised document text directly."
        )

        revised_content = llm_service.generate_chat_response(
            messages=[{"role": "user", "content": prompt}],
            system_instruction="You are an expert document editor. Apply requested changes accurately while maintaining tone and formatting."
        )

        base_name = orig_file.stem
        ext = orig_file.suffix.lower()
        new_filename = f"{base_name}_updated_{uuid.uuid4().hex[:6]}{ext}"
        new_output_path = settings.GENERATED_DIR / new_filename

        # Save revised version depending on format
        if ext in [".docx", ".doc"]:
            doc = DocxDocument()
            doc.add_heading(f"Revised: {base_name}", level=0)
            for line in revised_content.split("\n"):
                stripped = line.strip()
                if stripped.startswith("## "):
                    doc.add_heading(stripped[3:], level=2)
                elif stripped.startswith("# "):
                    doc.add_heading(stripped[2:], level=1)
                elif stripped.startswith("- ") or stripped.startswith("* "):
                    doc.add_paragraph(stripped[2:], style='List Bullet')
                elif stripped:
                    doc.add_paragraph(stripped)
            doc.save(str(new_output_path))
            
        elif ext == ".pdf":
            doc = SimpleDocTemplate(str(new_output_path), pagesize=letter)
            styles = getSampleStyleSheet()
            flowables = [Paragraph(f"Revised: {base_name}", styles['Title']), Spacer(1, 14)]
            for line in revised_content.split("\n"):
                stripped = line.strip()
                if stripped:
                    safe_line = stripped.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                    flowables.append(Paragraph(safe_line, styles['Normal']))
                    flowables.append(Spacer(1, 6))
            doc.build(flowables)
            
        else: # .txt, .md
            with open(str(new_output_path), "w", encoding="utf-8") as f:
                f.write(revised_content)

        summary = f"Updated document generated according to instructions: '{instructions}'."
        logger.info(f"Saved edited document to {new_output_path}. Original left untouched.")
        return new_filename, str(new_output_path), summary

document_service = DocumentService()
